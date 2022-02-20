import shutil

import numpy as np
import pandas as pd
from nltk.tokenize import sent_tokenize  # working with text
from docx import Document  # work with Word files
import xlsxwriter  # Excel
import os  # directory
import re  # check for punctuation
import string  # remove the punctuation

from database import db_utils


class TextDecompositionIntoSentencesAndWords:
    # find directory of the python program
    my_path = os.path.abspath(os.path.dirname(__file__))
    text_path = os.path.abspath(os.path.join(my_path, "../../", "Texts"))
    text_protected_path = os.path.abspath(os.path.join(my_path, "../../", "TextsFiltered–protected"))
    de_path = os.path.abspath(os.path.join(my_path, "../../", "Decomposed Texts"))
    all_decomposed_texts_path = os.path.join(de_path, 'All Decomposed Texts.xlsx')

    def decompose(self):
        for filename in os.listdir(self.de_path):
            os.remove(os.path.join(self.de_path, filename))

        # iterate over files in the directory
        for filename in os.listdir(self.text_path):
            # look only for Word files
            if filename.endswith(".docx") or filename.endswith(".doc"):

                file = str(os.path.join(self.text_path, filename))


                print("Now decomposing the text file: " + filename)

                # open the file
                text = Document(file)

                # remove .docx from the name
                # create a new Excel file with the name of the file
                decomposed_path = os.path.abspath(
                    os.path.join(self.my_path, "../../", "Decomposed Texts", filename[:-5] + '.xlsx'))
                workbook = xlsxwriter.Workbook(decomposed_path)

                worksheet_sentences = workbook.add_worksheet("Decomposition into sentences")

                worksheet_sentences.write('A1', "news ID")
                worksheet_sentences.write('B1', "contains sentence")
                worksheet_sentences.write('C1', "TextSentence")

                worksheet_words = workbook.add_worksheet("Decomposition into words")

                worksheet_words.write('A1', "news ID")
                worksheet_words.write('B1', "contains sentence ID")
                worksheet_words.write('C1', "word")
                worksheet_words.write('D1', "position")

                # to skip the first line that has "news ID" etc...
                sentences_row = 1
                sentenceID = 0
                words_row = 1
                # word = 0

                # iterate over the paragraphs
                for par in text.paragraphs:
                    para = par.text
                    list_of_sentences = sent_tokenize(para)

                    # iterate over the sentences
                    for i in range(len(list_of_sentences)):
                        worksheet_sentences.write(sentences_row, 0, os.path.splitext(filename)[0])
                        worksheet_sentences.write(sentences_row, 1, "s" + str(sentenceID).zfill(2))
                        worksheet_sentences.write(sentences_row, 2, list_of_sentences[i])
                        # words

                        word = 0

                        sentence = list_of_sentences[i].split()

                        # iterate over the words
                        for j in range(len(list_of_sentences[i].split())):
                            substring = '~'


                            if substring in sentence[j]:

                                new = ''
                                lt = list(sentence[j])
                                #print(lt)
                                for letter in lt:
                                    if (letter != '~'):
                                        if(letter != '“' and letter != '”'):
                                            new +=letter
                                    else:
                                        break




                                new.replace(' \' s', '').replace('’s', '').replace('’', '').replace('“', '').replace('”', '').replace('”','')
                                new = new.translate(str.maketrans('', '', string.punctuation))


                                # check this because some times there are only dashes as words
                                # no need to add them
                                if len(new) > 0:
                                    worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                    worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))

                                    new += '~'
                                    worksheet_words.write(words_row, 2, new)
                                    worksheet_words.write(words_row, 3, word)

                                    words_row += 1
                                    word += 1

                            else:
                                if (bool(re.match('^[a-zA-Z0-9]*$', sentence[j])) == True):
                                    # no punctuation
                                    # write the word sentence[j]
                                    if (sentence[j] == 'the' or sentence[j] == 'The'):
                                        t = list(sentence[j + 1])
                                        if (t[0] == 'o' or t[0] == 'u' or t[0] == 'e' or t[0] == 'i' or t[0] == 'y' or t[0] == 'a'):
                                            worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                            worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                            worksheet_words.write(words_row, 2,sentence[j].replace('“', '').replace('”', ''))
                                            worksheet_words.write(words_row, 3, word)

                                            words_row += 1
                                            word += 1

                                        else:
                                            sentence[j] +='\''
                                            worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                            worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                            worksheet_words.write(words_row, 2,sentence[j].replace('“', '').replace('”', ''))
                                            worksheet_words.write(words_row, 3, word)

                                            words_row += 1
                                            word += 1

                                    else:
                                        worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                        worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                        worksheet_words.write(words_row, 2, sentence[j].replace('“', '').replace('”', ''))
                                        worksheet_words.write(words_row, 3, word)

                                        words_row += 1
                                        word += 1

                                else:

                                    new = ''
                                    clean = list(sentence[j])
                                    # print(lt)
                                    for letter in clean:

                                        if (letter != '“' and letter != '”'):
                                            new += letter

                                    new.replace(' \' s', '').replace('’s', '').replace('’', '').replace('“','').replace('”', '').replace('”', '')
                                    new = new.translate(str.maketrans('', '', string.punctuation))



                                    # check this because some times there are only dashes as words
                                    # no need to add them
                                    if len(new) > 0:

                                        if (new == 'the' or new == 'The'):
                                            t = list(sentence[j + 1])
                                            if (t[0] == 'o' or t[0] == 'u' or t[0] == 'e' or t[0] == 'i' or t[
                                                0] == 'y' or t[0] == 'a'):
                                                worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                                worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                                worksheet_words.write(words_row, 2,
                                                                      new.replace('“', '').replace('”', ''))
                                                worksheet_words.write(words_row, 3, word)

                                                words_row += 1
                                                word += 1

                                            else:
                                                new += '\''
                                                worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                                worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                                worksheet_words.write(words_row, 2,
                                                                      new.replace('“', '').replace('”', ''))
                                                worksheet_words.write(words_row, 3, word)

                                                words_row += 1
                                                word += 1

                                        else:
                                            worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                                            worksheet_words.write(words_row, 1, "s" + str(sentenceID).zfill(2))
                                            worksheet_words.write(words_row, 2, new)
                                            worksheet_words.write(words_row, 3, word)

                                            words_row += 1
                                            word += 1

                        # write the whole sentence:
                        worksheet_words.write(words_row, 0, os.path.splitext(filename)[0])
                        worksheet_words.write(words_row, 1, "sentence " + str(sentenceID).zfill(2))
                        worksheet_words.write(words_row, 2, list_of_sentences[i])
                        words_row += 1
                        word += 1

                        sentences_row += 1
                        sentenceID += 1

                    # when the paragraph ends, add "0"
                    worksheet_sentences.write(sentences_row, 0, os.path.splitext(filename)[0])
                    worksheet_sentences.write(sentences_row, 1, "s" + str(sentenceID).zfill(2))
                    worksheet_sentences.write(sentences_row, 2, "0")
                    sentences_row += 1
                    sentenceID += 1

                # close the Excel file
                workbook.close()

    def merge(self):
        df_sentence = pd.DataFrame()
        df_words = pd.DataFrame()
        for filename in os.listdir(self.de_path):
            # look only for Excel files
            if filename.endswith(".xlsx"):
                file = str(os.path.join(self.de_path, filename))
                df_sentence = df_sentence.append(pd.read_excel(file, sheet_name="Decomposition into sentences"))
                df_words = df_words.append(pd.read_excel(file, sheet_name="Decomposition into words"))
                print("Files that were processed: " + file)

        df_words = df_words[~np.isnan(df_words['position'])]
        df_words = df_words[df_words['position'] != '']
        writer = pd.ExcelWriter(self.all_decomposed_texts_path)
        df_sentence.to_excel(writer, 'Decomposition into sentences')
        df_words.to_excel(writer, 'Decomposition into words')
        writer.save()

        words = pd.read_excel(self.all_decomposed_texts_path, sheet_name='Decomposition into words', usecols="B,C,D,E")

        conn = db_utils.db_connect()

        words.to_sql('words', con=conn, index=False, if_exists='append')
        conn.commit()

        sql = 'CREATE INDEX IF NOT EXISTS words_word_lower_idx ON words(lower(word))'
        cur = conn.cursor()
        cur.execute(sql)

        conn.commit()
        conn.close()

        print("The decomposed texts are fused in the file All Decomposed Texts.xlsx")

    def novels(self):
        conn = db_utils.db_connect()
        cur = conn.cursor()
        select_not_transcribed_words_sql = """
         SELECT W.WORD, COUNT(*) AS COUNT FROM WORDS W 
         WHERE NOT EXISTS 
         (SELECT 1 FROM DICTIONARY D WHERE LOWER(D.WORDFORM) = LOWER(W.WORD))
         AND NOT EXISTS (SELECT 1 FROM STOP_WORDS SW WHERE LOWER(SW." Stop Words") = LOWER(W.WORD))
         GROUP BY W.WORD"""
        cur.execute(select_not_transcribed_words_sql)
        my_result = cur.fetchall()
        cur.close()

        my_path = os.path.abspath(os.path.dirname(__file__))
        novel_path = os.path.abspath(os.path.join(my_path, '../../', 'Novel Word', 'Nontranscribed.xlsx'))
        pd.DataFrame(my_result).to_excel(novel_path, header=False, index=False)

        print("The words not found in the dictionary are ready in the folder Novel Words")

    def replace_with_protected_text(self):
        print("Start - Replace Texts with TextFilteredProtected")
        self.deleteTexts()
        for text_protected_file_name in os.listdir(self.text_protected_path):
            shutil.copy(os.path.join(self.text_protected_path, text_protected_file_name), self.text_path)

        print("End - Replace Texts with TextFilteredProtected")

    def deleteTexts(self):
        print("Start - Deleting Texts at: " + self.text_path)
        for text_file_name in os.listdir(self.text_path):
            os.remove(os.path.join(self.text_path, text_file_name))

        print("End - Deleting Texts at: " + self.text_path)
